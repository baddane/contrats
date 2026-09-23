"""Génère les dossiers de régularisation des contrats CIA (note ANAPEC/DPES/13/26).

Pour chaque contrat du fichier de suivi :
  - l'avenant de régularisation (FR) pré-rempli ;
  - le nouveau contrat de stage de formation-insertion (AR) pré-rempli ;
Les avenants sont numérotés 01/2026, 02/2026, … séparément pour chaque agence,
par ordre de date de signature du contrat initial.

Usage : python3 generer_regularisation.py
"""
import csv
import datetime as dt
import re
from pathlib import Path

import docx

ROOT = Path(__file__).parent
CSV = ROOT / "CIA-A-régulariser-Souss Massa.csv"
TPL_AVENANT = ROOT / "avenant regularisation 11_09.docx"
TPL_CONTRAT = ROOT / "نموذج اتفاقية-التدريب بقصد التكوين من أجل الادماج.docx"
OUT = ROOT / "regularisation"

# Agences traitées.
AGENCES = {"AGADIR", "INEZGANE AIT MELLOUL"}

RLM = "‏"
LEADER = re.compile(r"[.…]{2,}[.…  ]*")

AGENCE_AR = {
    "AGADIR": "أكادير",
    "AGADIR AGENCE UNIVERSITAIRE": "الوكالة الجامعية أكادير",
    "INEZGANE AIT MELLOUL": "إنزكان آيت ملول",
    "TAROUDANT": "تارودانت",
    "TIZNIT": "تيزنيت",
}
VILLE_FR = {
    "AGADIR": "Agadir",
    "AGADIR AGENCE UNIVERSITAIRE": "Agadir",
    "INEZGANE AIT MELLOUL": "Inezgane",
    "TAROUDANT": "Taroudant",
    "TIZNIT": "Tiznit",
}
VILLE_AR = {
    "AGADIR": "أكادير",
    "AGADIR AGENCE UNIVERSITAIRE": "أكادير",
    "INEZGANE AIT MELLOUL": "إنزكان",
    "TAROUDANT": "تارودانت",
    "TIZNIT": "تيزنيت",
}
SECTEUR_AR = {
    "Activités immobilières": "الأنشطة العقارية",
    "Agriculture, chasse": "الفلاحة والقنص",
    "Autres": "أنشطة أخرى",
    "Commerce de détail et réparation d'articles domestiques": "تجارة التقسيط وإصلاح الأدوات المنزلية",
    "Commerce de gros et intermédiaires du commerce": "تجارة الجملة والوساطة التجارية",
    "Conseil en systèmes informatiques": "الاستشارة في الأنظمة المعلوماتية",
    "Construction": "البناء",
    "Fabrication de machines et appareils électriques": "صناعة الآلات والأجهزة الكهربائية",
    "Fabrication de machines et équipements": "صناعة الآلات والمعدات",
    "Hôtellerie et restauration": "الفندقة والمطعمة",
    "Intermédiation financière": "الوساطة المالية",
    "Location sans opérateur": "الكراء بدون مشغل",
    "Santé et action sociale": "الصحة والعمل الاجتماعي",
    "Services auxiliaires des transports": "الخدمات المساعدة للنقل",
    "Services fournis principalement aux entreprises": "الخدمات المقدمة أساسا للمقاولات",
    "Services personnels": "الخدمات الشخصية",
    "Transports terrestres": "النقل البري",
}
METIER_AR = {
    "Animatrice de mariage": "تنشيط حفلات الزفاف",
    "Commercial/Commerciale": "مندوب(ة) تجاري(ة)",
    "Conducteur de transport en commun": "سياقة وسائل النقل العمومي",
    "Electricien de maintenance": "كهربائي(ة) الصيانة",
    "Employé polyvalent de restauration": "مستخدم(ة) متعدد(ة) المهام في المطعمة",
    "Magasinier d'entrepôt": "أمين(ة) مخزن",
    "Mécanicien d'entretien sur machines industrielles": "ميكانيكي(ة) صيانة الآلات الصناعية",
    "Responsable de la gestion des ressources humaines": "مسؤول(ة) تدبير الموارد البشرية",
    "Vendeur/Vendeuse": "بائع(ة)",
}
NATIONALITE_AR = {"0": "مغربية"}
# Conseiller ANAPEC signataire, par agence.
CONSEILLER = {
    "AGADIR": "BRAHIM ALAAOUCH",
    "INEZGANE AIT MELLOUL": "SAID EN KHAL",
}

def fdate(s):
    """'8/13/2026' (format US du fichier) -> '13/08/2026'."""
    return dt.datetime.strptime(s.strip(), "%m/%d/%Y").strftime("%d/%m/%Y") if s.strip() else ""


def clean(s):
    return " ".join(s.split())


def replace_span(p, start, end, value):
    """Remplace les caractères [start, end) du texte du paragraphe, en gardant la mise en forme des runs."""
    pos, placed = 0, False
    for run in p.runs:
        t = run.text
        a, b = pos, pos + len(t)
        pos = b
        if b <= start or a >= end:
            continue
        s, e = max(start, a) - a, min(end, b) - a
        run.text = t[:s] + ("" if placed else value) + t[e:]
        placed = True


def fill(p, label, value, rtl=False):
    """Remplace la première ligne de pointillés qui suit `label` dans le paragraphe."""
    if not value:
        return
    text = p.text
    i = text.find(label)
    if i < 0:
        raise ValueError(f"Libellé introuvable : {label!r} dans {text[:60]!r}")
    m = LEADER.search(text, i + len(label))
    if not m:
        raise ValueError(f"Pointillés introuvables après {label!r}")
    value = f" {value} "
    if rtl:
        value = f"{RLM}{value}{RLM}"
    replace_span(p, m.start(), m.end(), value)


def _insert(p, index, value):
    """Insère `value` à la position `index` du texte du paragraphe."""
    pos = 0
    for run in p.runs:
        t = run.text
        if index <= pos + len(t):
            k = index - pos
            run.text = t[:k] + value + t[k:]
            return
        pos += len(t)


def para(doc, contains):
    for p in doc.paragraphs:
        if contains in p.text:
            return p
    raise ValueError(f"Paragraphe introuvable : {contains!r}")


def avenant(r, numero):
    d = docx.Document(TPL_AVENANT)
    ref = r["REF_CONTRAT"].strip()
    nom = clean(f"{r['NOM_CANDIDAT']} {r['PRENOM']}").upper()

    fill(para(d, "AVENANT DE RÉGULARISATION"), "N°", numero)
    p = para(d, "Raison Sociale")
    fill(p, "Raison Sociale", clean(r["RAISON_SOCIALE"]))
    fill(p, "Adresse :", clean(r["ADRESSE"]))
    p = para(d, "Nom et prénom")
    fill(p, "Nom et prénom :", nom)
    fill(p, "CIN :", r["CIN"].strip().upper())
    fill(p, "Date de naissance :", fdate(r["DATE_NAISSANCE"]))
    agence = r["NOM_AGENCE"].strip()
    fill(para(d, "représentée par l’agence"), "l’agence",
         f"ANAPEC {agence}, en la personne de M. {CONSEILLER[agence]}")
    p = para(d, "ci-après dénommé le « contrat initial »")
    fill(p, "n°", ref)
    fill(p, "signé le", fdate(r["DATE_SIGNATURE"]))
    fill(para(d, "portant la même référence"), "référence n°", ref)
    fill(para(d, "soit le"), "soit le", fdate(r["DATE_EFFET"]) + ".")
    p = para(d, "Fait à")
    fill(p, "Fait à", VILLE_FR[agence])
    fill(p, ", le", fdate(r["DATE_SIGNATURE"]))
    return d


def contrat(r):
    d = docx.Document(TPL_CONTRAT)
    agence = r["NOM_AGENCE"].strip()
    nom = clean(f"{r['NOM_CANDIDAT']} {r['PRENOM']}").upper()
    secteur = clean(r["SECTEUR_ACTIVITE"])
    metier = clean(r["EMPLOI_METIER"])

    p = para(d, "الوكالة :")
    _insert(p, len(p.text), AGENCE_AR[agence])
    fill(para(d, "الاتفاقية رقم"), "الاتفاقية رقم", r["REF_CONTRAT"].strip(), rtl=True)
    fill(para(d, "الاسم أو العنوان التجاري"), "العنوان التجاري", clean(r["RAISON_SOCIALE"]), rtl=True)
    fill(para(d, "قطاع النشاط"), "قطاع النشاط", SECTEUR_AR.get(secteur, secteur))
    fill(para(d, "العنوان.."), "العنوان", clean(r["ADRESSE"]), rtl=True)
    fill(para(d, "الهاتف + الفاكس"), "الهاتف + الفاكس", r["TEL1"].strip(), rtl=True)
    fill(para(d, "رقم الانخراط في الصندوق"), "للضمان الاجتماعي", r["NUM_CNSS"].strip(), rtl=True)

    fill(para(d, "الاسم العائلي والشخصي"), "والشخصي", nom, rtl=True)
    nat = r["NATIONALITE"].strip()
    fill(para(d, "الجنسية"), "الجنسية", NATIONALITE_AR.get(nat, nat))
    fill(para(d, "رقم البطاقة الوطنية"), "بطاقة الإقامة.", r["CIN"].strip().upper(), rtl=True)
    fill(para(d, "رقم التسجيل بالصندو"), "الاجتماعي:", r["NUM_CNSS_1"].strip(), rtl=True)
    p = para(d, "غير حاصل على شهادة")
    if r["CATEGORIE_CONTRAT"].strip() == "avec_diplome":
        _insert(p, p.text.index("نعم،") + len("نعم،"), " ✔")
    else:
        _insert(p, p.text.index("غير حاصل على شهادة"), "✔ ")

    fill(para(d, "بتدريب السيد"), "السيدة", nom, rtl=True)
    fill(para(d, "شهرا (12"), "لمدة", "12")
    if metier:
        fill(para(d, "للقيام بأنشطة"), "أنشطة", METIER_AR.get(metier, metier))
    fill(para(d, "يحدد مبلغها"), "في", r["REMUNERATION"].strip())
    p = para(d, "حرر ب")
    fill(p, "حرر ب", VILLE_AR[agence])
    fill(p, "بتاريخ", fdate(r["DATE_SIGNATURE"]), rtl=True)
    d.tables[0].cell(1, 0).paragraphs[0].add_run(CONSEILLER[agence])
    return d


def safe(s):
    return re.sub(r'[\\/:*?"<>|]+', "-", s.strip())


def main():
    with open(CSV, encoding="cp1252", newline="") as f:
        rows = [r for r in csv.DictReader(f) if r["NOM_AGENCE"].strip() in AGENCES]

    rows.sort(key=lambda r: (dt.datetime.strptime(r["DATE_SIGNATURE"].strip(), "%m/%d/%Y"),
                             r["REF_CONTRAT"].strip()))
    compteurs = {}
    for r in rows:
        agence = r["NOM_AGENCE"].strip()
        compteurs[agence] = compteurs.get(agence, 0) + 1
        numero = f"{compteurs[agence]:02d}/2026"
        ref = safe(r["REF_CONTRAT"])
        nom = safe(clean(f"{r['NOM_CANDIDAT']} {r['PRENOM']}").upper())
        dossier = OUT / safe(r["NOM_AGENCE"]) / f"{ref} - {nom}"
        dossier.mkdir(parents=True, exist_ok=True)
        contrat(r).save(dossier / f"1 - Nouveau contrat {ref}.docx")
        avenant(r, numero).save(dossier / f"2 - Avenant N° {safe(numero)}.docx")

    print(f"{len(rows)} dossiers générés dans {OUT}")


if __name__ == "__main__":
    main()
